from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from modules.template.schema_extractor_xlsx import extract_xlsx_schema


def test_extract_xlsx_schema_reads_named_ranges(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "placeholder"
    wb.defined_names.add(DefinedName(name="project_name", attr_text="'Sheet1'!$A$1"))
    wb.save(path)
    schema = extract_xlsx_schema(path)
    assert schema.source_format == "xlsx"
    assert any(p.placeholder_id == "project_name" for p in schema.placeholders)

